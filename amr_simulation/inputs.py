"""Workload caching, layout normalization, assignment, and preflight checks."""

from __future__ import annotations

import hashlib
import json
import math
import os
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from warehouse_layout.domain import GridProject

from .models import (
    Rack,
    SimulationConfig,
    ValidationReport,
    Workload,
    WorkloadTask,
    grid_name,
    grid_position,
)
from .routing import GridRouter


CACHE_SCHEMA = "amr_workload_cache/v1"
REQUIRED_COLUMNS = ("Date", "Store ID", "Item or SKU")
DEFAULT_CACHE_DIR = Path(__file__).resolve().parent / ".cache"


def _date_value(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    try:
        return datetime.fromisoformat(str(value).strip()).date()
    except (TypeError, ValueError):
        return None


def _identifier(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _fingerprint(path: Path) -> dict:
    resolved = path.expanduser().resolve()
    stat = resolved.stat()
    return {
        "source_path": str(resolved),
        "source_size": stat.st_size,
        "source_mtime_ns": stat.st_mtime_ns,
    }


def _cache_path(source: Path, cache_dir: Path) -> Path:
    digest = hashlib.sha256(str(source.resolve()).encode()).hexdigest()[:20]
    return cache_dir / f"{digest}.npz"


def _metadata_array(value: dict) -> np.ndarray:
    return np.frombuffer(json.dumps(value, separators=(",", ":")).encode(), dtype=np.uint8)


def _decode_metadata(value: np.ndarray) -> dict:
    return json.loads(value.tobytes().decode())


def _load_cached_workload(source: Path, cache_dir: Path, fingerprint: dict) -> Workload | None:
    path = _cache_path(source, cache_dir)
    if not path.exists():
        return None
    try:
        with np.load(path, allow_pickle=False) as cached:
            metadata = _decode_metadata(cached["metadata"])
            if metadata.get("schema") != CACHE_SCHEMA:
                return None
            if any(metadata.get(key) != value for key, value in fingerprint.items()):
                return None
            dates = cached["task_dates"]
            releases = cached["release_seconds"]
            stores = cached["store_indices"]
            offsets = cached["sku_offsets"]
            sku_indices = cached["sku_indices"]
            line_counts = cached["line_counts"]
            if not (len(dates) == len(releases) == len(stores) == len(offsets) - 1):
                raise ValueError("cache task arrays have inconsistent lengths")
            sku_names, store_names = metadata["skus"], metadata["stores"]
            tasks = []
            for index, ordinal in enumerate(dates):
                first, last = int(offsets[index]), int(offsets[index + 1])
                counts = {
                    sku_names[int(sku_indices[item])]: int(line_counts[item])
                    for item in range(first, last)
                }
                task_date = date.fromordinal(int(ordinal))
                store = store_names[int(stores[index])]
                tasks.append(
                    WorkloadTask(
                        f"{task_date.isoformat()}/{store}", task_date,
                        0.0, store, counts,
                    )
                )
            return Workload(
                tasks, date.fromisoformat(metadata["min_date"]),
                date.fromisoformat(metadata["max_date"]), source,
                int(metadata["valid_rows"]), True,
            )
    except Exception:
        try:
            path.unlink()
        except OSError:
            pass
        return None


def _find_sheet(workbook):
    required = set(REQUIRED_COLUMNS)
    for worksheet in workbook.worksheets:
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(50, worksheet.max_row), values_only=True), 1
        ):
            positions = {
                str(value).strip(): index
                for index, value in enumerate(values)
                if value is not None
            }
            if required.issubset(positions):
                return worksheet, row_number, positions
    raise ValueError("no worksheet contains required columns: " + ", ".join(REQUIRED_COLUMNS))


def load_workload(path: Path, cache_dir: Path | None = None, use_cache: bool = True) -> Workload:
    source = Path(path).expanduser().resolve()
    if not source.exists():
        raise ValueError(f"order workbook not found: {source}")
    cache_root = Path(cache_dir or DEFAULT_CACHE_DIR)
    fingerprint = _fingerprint(source)
    if use_cache:
        cached = _load_cached_workload(source, cache_root, fingerprint)
        if cached is not None:
            return cached

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet, header_row, positions = _find_sheet(workbook)
        grouped: dict[tuple[int, str], Counter[str]] = {}
        invalid_count, invalid_samples = 0, []
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
        ):
            picked_date = _date_value(values[positions["Date"]] if positions["Date"] < len(values) else None)
            store = _identifier(values[positions["Store ID"]] if positions["Store ID"] < len(values) else None)
            sku = _identifier(values[positions["Item or SKU"]] if positions["Item or SKU"] < len(values) else None)
            if picked_date is None or not store or not sku:
                invalid_count += 1
                if len(invalid_samples) < 20:
                    invalid_samples.append(row_number)
                continue
            key = picked_date.toordinal(), store
            grouped.setdefault(key, Counter())[sku] += 1
        if invalid_count:
            raise ValueError(
                f"workbook contains {invalid_count:,} invalid required rows; "
                f"sample row numbers: {invalid_samples}"
            )
        if not grouped:
            raise ValueError("workbook contains no valid order rows")
    finally:
        workbook.close()

    tasks = [
        WorkloadTask(
            f"{date.fromordinal(ordinal).isoformat()}/{store}", date.fromordinal(ordinal),
            0.0, store, dict(sorted(counts.items())),
        )
        for (ordinal, store), counts in sorted(grouped.items())
    ]
    workload = Workload(
        tasks, tasks[0].task_date, tasks[-1].task_date, source,
        sum(task.source_lines for task in tasks), False,
    )
    if use_cache:
        _write_workload_cache(workload, cache_root, fingerprint, worksheet.title)
    return workload


def _write_workload_cache(
    workload: Workload, cache_dir: Path, fingerprint: dict, worksheet: str
) -> None:
    cache_dir.mkdir(parents=True, exist_ok=True)
    skus = sorted({sku for task in workload.tasks for sku in task.line_counts})
    stores = sorted({task.store_id for task in workload.tasks})
    sku_lookup, store_lookup = {v: i for i, v in enumerate(skus)}, {v: i for i, v in enumerate(stores)}
    offsets, sku_indices, line_counts = [0], [], []
    for task in workload.tasks:
        for sku, count in task.line_counts.items():
            sku_indices.append(sku_lookup[sku])
            line_counts.append(count)
        offsets.append(len(sku_indices))
    metadata = {
        "schema": CACHE_SCHEMA, **fingerprint, "worksheet": worksheet,
        "skus": skus, "stores": stores, "valid_rows": workload.valid_rows,
        "min_date": workload.min_date.isoformat(), "max_date": workload.max_date.isoformat(),
    }
    path = _cache_path(workload.source_path, cache_dir)
    temporary = path.with_suffix(".tmp")
    try:
        with temporary.open("wb") as stream:
            np.savez_compressed(
                stream, metadata=_metadata_array(metadata),
                task_dates=np.array([task.task_date.toordinal() for task in workload.tasks], dtype=np.int32),
                release_seconds=np.array([task.release_seconds for task in workload.tasks], dtype=np.float64),
                store_indices=np.array([store_lookup[task.store_id] for task in workload.tasks], dtype=np.int32),
                sku_offsets=np.array(offsets, dtype=np.int64),
                sku_indices=np.array(sku_indices, dtype=np.int32),
                line_counts=np.array(line_counts, dtype=np.int32),
            )
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def racks_from_layout(assignments: list[dict]) -> dict[str, Rack]:
    rack_skus: dict[str, set[str]] = defaultdict(set)
    for row in assignments:
        if row.get("assignment_status") not in (None, "ASSIGNED"):
            continue
        sku = _identifier(row.get("sku"))
        rack_id = _identifier(row.get("rack_id"))
        if sku and rack_id:
            rack_skus[rack_id].add(sku)
    return {
        rack_id: Rack(rack_id, grid_position(rack_id), frozenset(skus))
        for rack_id, skus in sorted(rack_skus.items())
    }


def assign_workstations(
    tasks: list[WorkloadTask], config: SimulationConfig
) -> dict[str, str]:
    workloads = Counter()
    for task in tasks:
        workloads[task.store_id] += task.source_lines
    unknown = sorted(set(config.store_workstation_overrides) - set(workloads))
    if unknown:
        raise ValueError("workstation overrides reference unknown stores: " + ", ".join(unknown))
    loads = {station: 0 for station in config.workstations}
    mapping: dict[str, str] = {}
    for store, station in sorted(config.store_workstation_overrides.items()):
        if station not in loads:
            raise ValueError(f"override for {store} references disabled workstation {station}")
        mapping[store] = station
        loads[station] += workloads[store]
    remaining = sorted(
        (store for store in workloads if store not in mapping),
        key=lambda store: (-workloads[store], store),
    )
    for store in remaining:
        station = min(loads, key=lambda value: (loads[value], value))
        mapping[store] = station
        loads[station] += workloads[store]
    return mapping


def validate_inputs(
    project: GridProject,
    router: GridRouter,
    racks: dict[str, Rack],
    tasks: list[WorkloadTask],
    config: SimulationConfig,
    workstation_mapping: dict[str, str],
) -> ValidationReport:
    report = ValidationReport()
    marker_racks = {grid_name(position): position for position in router.rack_positions}
    marker_stations = router.workstations
    for station in config.workstations:
        if station not in marker_stations:
            report.error("invalid_workstation", f"workstation {station} is absent from the grid", workstation=station)
    for store, station in workstation_mapping.items():
        if station not in config.workstations:
            report.error("invalid_store_mapping", f"store {store} maps to disabled workstation {station}", store_id=store, workstation=station)

    spawn_positions = []
    for node in config.spawn_nodes:
        position = grid_position(node)
        spawn_positions.append(position)
        if node not in marker_racks:
            report.error("invalid_spawn", f"spawn node {node} is not an active rack marker", node=node)

    absent_racks = sorted(set(racks) - set(marker_racks))
    for rack_id in absent_racks:
        report.error("layout_rack_absent", f"layout rack {rack_id} is absent from the grid", rack_id=rack_id, skus=sorted(racks[rack_id].skus))

    demand = Counter()
    for task in tasks:
        demand.update(task.line_counts)
    mapped_skus = {sku for rack in racks.values() for sku in rack.skus}
    for sku in sorted(set(demand) - mapped_skus):
        report.error("unmapped_sku", f"workload SKU {sku} is not assigned to a rack", sku=sku, historical_line_count=demand[sku])

    active_stations = [marker_stations[value] for value in config.workstations if value in marker_stations]
    for station_id in config.workstations:
        if station_id not in marker_stations:
            continue
        station = marker_stations[station_id]
        if not any(
            router.reachable(spawn, station)
            and router.reachable(station, spawn)
            for spawn in spawn_positions
        ):
            report.error(
                "disconnected_workstation",
                f"workstation {station_id} has no spawn-node round trip",
                workstation=station_id,
            )
    for rack_id, rack in racks.items():
        if rack_id in absent_racks:
            continue
        pickup_ok = any(router.reachable(spawn, rack.position) for spawn in spawn_positions)
        failed_stations = [
            station
            for station in config.workstations
            if station in marker_stations
            and (
                not router.reachable(rack.position, marker_stations[station])
                or not router.reachable(marker_stations[station], rack.position)
            )
        ]
        if not pickup_ok or failed_stations:
            line_count = sum(demand[sku] for sku in rack.skus)
            report.error(
                "unreachable_rack", f"occupied rack {rack_id} is unreachable under the rack-obstacle rule",
                rack_id=rack_id, skus=sorted(rack.skus), historical_line_count=line_count,
                pickup_reachable=pickup_ok, failed_workstations=failed_stations,
            )
    for node, spawn in zip(config.spawn_nodes, spawn_positions):
        if active_stations and not any(
            router.reachable(spawn, station) and router.reachable(station, spawn)
            for station in active_stations
        ):
            report.error("disconnected_spawn", f"spawn node {node} has no workstation round trip", node=node)
    report.warning(
        "reservation_model",
        "node ownership prevents overlap but does not use time-expanded reservations",
    )
    return report
