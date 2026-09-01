from __future__ import annotations

import json
import math
import gzip
import hashlib
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from .models import Grid, Node, Task


def _node(raw: dict) -> Node:
    return int(raw["column"]), int(raw["row"])


def load_grid(path: Path) -> Grid:
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not str(raw.get("schema", "")).startswith("rmf_grid_map_editor/"):
        raise ValueError("unsupported grid-project schema")
    spec = raw["grid"]
    sx, sy = float(spec["spacing_m"]), float(spec.get("spacing_y_m", spec["spacing_m"]))
    width = int(round(float(spec["width_m"]) / sx)) + 1
    height = int(round(float(spec["length_m"]) / sy)) + 1
    deleted = {_node(v) for v in raw.get("deleted_positions", [])}
    nodes = {(x, y) for x in range(width) for y in range(height)} - deleted
    deleted_lanes = {(_node(v["start"]), _node(v["end"])) for v in raw.get("deleted_lanes", [])}
    deleted_lanes |= {(b, a) for a, b in tuple(deleted_lanes)}
    one_way = {(_node(v["start"]), _node(v["end"])) for v in raw.get("one_way_lanes", [])}
    one_way_pairs = {frozenset((a, b)) for a, b in one_way}
    edges: set[tuple[Node, Node]] = set()
    for x, y in nodes:
        for nxt in ((x + 1, y), (x - 1, y), (x, y + 1), (x, y - 1)):
            if nxt not in nodes or ((x, y), nxt) in deleted_lanes:
                continue
            pair = frozenset(((x, y), nxt))
            if pair in one_way_pairs and ((x, y), nxt) not in one_way:
                continue
            edges.add(((x, y), nxt))
    racks, stations = {}, {}
    for marker in raw.get("markers", []):
        pos, endpoint = _node(marker), str(marker.get("endpoint_id", ""))
        if pos not in nodes or not endpoint:
            continue
        target = racks if marker.get("role") == "rack" else stations if marker.get("role") == "workstation" else None
        if target is not None:
            if endpoint in target:
                raise ValueError(f"duplicate endpoint {endpoint}")
            target[endpoint] = pos
    return Grid(width, height, sx, sy, nodes, edges, racks, stations)


def load_layout(path: Path) -> dict[str, set[str]]:
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not str(raw.get("schema", "")).startswith("inventory_slotting_layout/"):
        raise ValueError("unsupported slotting-layout schema")
    racks: dict[str, set[str]] = {}
    for row in raw.get("assignments", []):
        if row.get("assignment_status", "ASSIGNED") != "ASSIGNED":
            continue
        sku = str(row.get("sku", "")).strip()
        rack = str(row.get("rack_waypoint") or row.get("rack_id") or "").strip()
        if sku and rack:
            racks.setdefault(rack, set()).add(sku)
    if not racks:
        raise ValueError("layout contains no assigned SKUs")
    return racks


def _date(value) -> date | None:
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    try: return datetime.fromisoformat(str(value).strip()).date()
    except (TypeError, ValueError): return None


def _identifier(value) -> str:
    if value in (None, ""): return ""
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip()


def load_orders(path: Path) -> list[Task]:
    source = path.resolve(); stat = source.stat()
    cache_dir = Path(__file__).resolve().parents[1] / ".cache"; cache_dir.mkdir(exist_ok=True)
    digest = hashlib.sha256(str(source).encode()).hexdigest()[:20]
    cache = cache_dir / f"orders-{digest}.json.gz"
    if cache.exists():
        try:
            with gzip.open(cache, "rt", encoding="utf-8") as stream: saved = json.load(stream)
            if saved.get("schema") == "lacam_orders/v2" and saved["size"] == stat.st_size and saved["mtime_ns"] == stat.st_mtime_ns:
                return [Task(row["task_id"], date.fromisoformat(row["date"]), row["store"], row["release_seconds"], row["lines"]) for row in saved["tasks"]]
        except (OSError, KeyError, ValueError, json.JSONDecodeError): pass
    book = load_workbook(path, read_only=True, data_only=True)
    required = {"Date", "Store ID", "Item or SKU"}
    try:
        found = None
        for sheet in book.worksheets:
            for number, values in enumerate(sheet.iter_rows(max_row=min(50, sheet.max_row), values_only=True), 1):
                columns = {str(v).strip(): i for i, v in enumerate(values) if v is not None}
                if required <= columns.keys(): found = sheet, number, columns; break
            if found: break
        if not found: raise ValueError("order workbook lacks Date, Store ID, or Item or SKU")
        sheet, header, columns = found
        grouped: dict[tuple[date, str], Counter[str]] = {}
        for values in sheet.iter_rows(min_row=header + 1, values_only=True):
            when = _date(values[columns["Date"]] if columns["Date"] < len(values) else None)
            store = _identifier(values[columns["Store ID"]]) if columns["Store ID"] < len(values) else ""
            sku = _identifier(values[columns["Item or SKU"]]) if columns["Item or SKU"] < len(values) else ""
            if not when or not store or not sku: continue
            grouped.setdefault((when, store), Counter())[sku] += 1
    finally:
        book.close()
    tasks = [Task(f"{d.isoformat()}/{s}", d, s, 0.0, dict(lines))
             for (d, s), lines in sorted(grouped.items())]
    with gzip.open(cache, "wt", encoding="utf-8") as stream:
        json.dump({"schema": "lacam_orders/v2", "size": stat.st_size, "mtime_ns": stat.st_mtime_ns,
                   "tasks": [{"task_id": t.task_id, "date": t.date.isoformat(), "store": t.store,
                              "release_seconds": t.release_seconds, "lines": t.lines} for t in tasks]}, stream)
    return tasks


def assign_workstations(tasks: list[Task], stations: tuple[str, ...]) -> dict[str, str]:
    demand = Counter()
    for task in tasks: demand[task.store] += sum(task.lines.values())
    loads = {station: 0 for station in stations}
    result = {}
    for store in sorted(demand, key=lambda s: (-demand[s], s)):
        station = min(stations, key=lambda s: (loads[s], s))
        result[store] = station; loads[station] += demand[store]
    return result


def validate(grid: Grid, layout: dict[str, set[str]], tasks: list[Task], spawns: list[Node], stations: tuple[str, ...]) -> dict:
    errors, warnings = [], []
    for node in spawns:
        if node not in grid.nodes: errors.append({"code": "spawn_absent", "node": node})
    for station in stations:
        if station not in grid.workstations: errors.append({"code": "station_absent", "station": station})
    for rack in layout:
        if rack not in {f"G{x}_{y}" for x, y in grid.nodes}: errors.append({"code": "rack_absent", "rack": rack})
    available = set().union(*layout.values())
    missing = sorted({sku for task in tasks for sku in task.lines} - available)
    if missing: errors.append({"code": "missing_skus", "count": len(missing), "sample": missing[:20]})
    return {"valid": not errors, "errors": errors, "warnings": warnings}
