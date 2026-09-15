from collections import Counter, defaultdict
from pathlib import Path
import csv
from openpyxl import load_workbook

SOURCE = Path("resources/data/Sample Data.xlsx")
OUTPUT = Path("resources/data/mixed_uom_order_lines.csv")

workbook = load_workbook(SOURCE, read_only=True, data_only=True)
worksheet = workbook["Picking"]
rows = worksheet.iter_rows(min_row=4, values_only=True)
headers = next(rows)
positions = {str(value).strip(): i for i, value in enumerate(headers) if value is not None}
sku_index = positions["Item or SKU"]
uom_index = positions["UOM"]
conversion_index = positions["UOM ConversionQty"]
records = []
variants = defaultdict(set)

for values in rows:
    sku = "" if values[sku_index] is None else str(values[sku_index]).strip()
    if not sku:
        continue
    uom = "" if values[uom_index] is None else str(values[uom_index]).strip()
    conversion = values[conversion_index]
    variants[sku].add((uom, str(conversion)))
    records.append(values)

mixed = {sku for sku, values in variants.items() if len(values) > 1}
fields = [
    "Date", "Time", "Store ID", "Item or SKU", "Quantity (in EA)",
    "Quantity UOM", "UOM", "UOM Adj.", "UOM ConversionQty", "Length (m)",
    "Width (m)", "Height (m)", "Weight (kg)", "Size",
]
fields.append("UOM status")
records = [
    values for values in records
    if str(values[sku_index]).strip() in mixed
]
variant_counts = {
    sku: Counter((str(values[uom_index]).strip(), str(values[conversion_index]))
                 for values in records if str(values[sku_index]).strip() == sku)
    for sku in mixed
}
records.sort(key=lambda values: (str(values[sku_index]).strip(), values[0] or "", values[1] or "", str(values[2] or "")))
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
with OUTPUT.open("w", newline="", encoding="utf-8") as stream:
    writer = csv.writer(stream)
    writer.writerow(fields)
    for values in records:
        sku = "" if values[sku_index] is None else str(values[sku_index]).strip()
        variant = (str(values[uom_index]).strip(), str(values[conversion_index]))
        status = "COMMON" if variant_counts[sku][variant] == max(variant_counts[sku].values()) else "UNCOMMON"
        writer.writerow([*values[:len(fields) - 1], status])
workbook.close()
print(f"{len(mixed)} mixed-UOM SKUs; output {OUTPUT}")
