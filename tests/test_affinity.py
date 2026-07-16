"""Tests for line-order SKU/store affinity analysis."""

from __future__ import annotations

import csv
import json
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

from warehouse_layout.affinity import (
    AffinityCancelledError,
    AffinityService,
    EXPORT_SCHEMA,
)


class AffinityServiceTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.workbook_path = self.root / "orders.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Picking"
        sheet.append(["Report heading"])
        sheet.append(["Date", "Store ID", "Item or SKU", "Quantity (in EA)"])
        sheet.append([datetime(2026, 1, 1), "S1", "A", 1])
        sheet.append([datetime(2026, 1, 1), "S1", "A", 999])
        sheet.append([datetime(2026, 1, 1), "S2", "A", 4])
        sheet.append([datetime(2026, 1, 1), "S1", "B", 2])
        sheet.append([datetime(2026, 1, 2), "S2", "B", 2])
        sheet.append([datetime(2026, 1, 2), "S3", "C", 1])
        sheet.append(["not-a-date", "S1", "A", 1])
        sheet.append([datetime(2026, 1, 2), "", "A", 1])
        workbook.save(self.workbook_path)
        self.service = AffinityService(self.root / "cache")

    def tearDown(self):
        self.temporary.cleanup()

    def test_line_counts_and_cosine_affinity_ignore_quantity(self):
        dataset = self.service.load_orders(self.workbook_path)
        analysis = self.service.analyze(dataset)
        a = dataset.skus.index("A")
        b = dataset.skus.index("B")
        s1 = dataset.stores.index("S1")
        s2 = dataset.stores.index("S2")
        self.assertEqual(dataset.valid_rows, 6)
        self.assertEqual(dataset.skipped_rows, 2)
        self.assertEqual(int(analysis.frequency[a, s1]), 2)
        self.assertEqual(int(analysis.frequency[a, s2]), 1)
        self.assertEqual(int(analysis.sku_totals[a]), 3)
        self.assertEqual(int(analysis.shared_store_days[a, b]), 1)
        self.assertEqual(int(analysis.sku_store_day_totals[a]), 2)
        self.assertEqual(int(analysis.sku_store_day_totals[b]), 2)
        self.assertAlmostEqual(float(analysis.similarity[a, b]), 0.5)

    def test_date_filter_is_inclusive_and_empty_range_is_rejected(self):
        dataset = self.service.load_orders(self.workbook_path)
        analysis = self.service.analyze(
            dataset, date(2026, 1, 2), date(2026, 1, 2)
        )
        self.assertEqual(analysis.event_count, 2)
        self.assertEqual(set(analysis.dataset.skus[index] for index in analysis.active_sku_indices()), {"B", "C"})
        with self.assertRaisesRegex(ValueError, "No valid order events"):
            self.service.analyze(dataset, date(2025, 1, 1), date(2025, 1, 2))
        with self.assertRaisesRegex(ValueError, "Start date"):
            self.service.analyze(dataset, date(2026, 1, 2), date(2026, 1, 1))

    def test_cache_reuse_and_corrupt_cache_recovery(self):
        first = self.service.load_orders(self.workbook_path)
        self.assertFalse(first.cache_used)
        second = self.service.load_orders(self.workbook_path)
        self.assertTrue(second.cache_used)
        cache_path = self.service._cache_path(self.workbook_path)
        cache_path.write_bytes(b"not a numpy archive")
        recovered = self.service.load_orders(self.workbook_path)
        self.assertFalse(recovered.cache_used)
        self.assertEqual(recovered.valid_rows, first.valid_rows)
        self.assertTrue(cache_path.exists())

    def test_source_change_invalidates_cache(self):
        self.service.load_orders(self.workbook_path)
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Date", "Store ID", "Item or SKU"])
        sheet.append([datetime(2026, 2, 1), "S9", "Z",])
        workbook.save(self.workbook_path)
        changed = self.service.load_orders(self.workbook_path)
        self.assertFalse(changed.cache_used)
        self.assertEqual(changed.skus, ("Z",))

    def test_missing_headers_and_cancellation_are_reported(self):
        missing = self.root / "missing.xlsx"
        workbook = Workbook()
        workbook.active.append(["Date", "SKU"])
        workbook.save(missing)
        with self.assertRaisesRegex(ValueError, "required columns"):
            self.service.load_orders(missing, use_cache=False)
        with self.assertRaises(AffinityCancelledError):
            self.service.load_orders(
                self.workbook_path, cancelled=lambda: True
            )
        self.assertFalse(self.service._cache_path(self.workbook_path).exists())

    def test_relationship_threshold_can_leave_sku_without_matches(self):
        dataset = self.service.load_orders(self.workbook_path)
        analysis = self.service.analyze(dataset)
        c = dataset.skus.index("C")
        self.assertEqual(analysis.related_skus(c, min_shared_store_days=2), [])

    def test_slotting_thresholds_are_derived_again_for_each_workbook(self):
        def build(path, scale):
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Date", "Store ID", "Item or SKU"])
            picked_date = date(2025, 1, 1)
            for store, pair, count in (
                ("S_AB", ("A", "B"), 5),
                ("S_AC", ("A", "C"), 2),
                ("S_BC", ("B", "C"), 1),
            ):
                for _ in range(count * scale):
                    for sku in pair:
                        sheet.append([picked_date, store, sku])
                    picked_date += timedelta(days=1)
            workbook.save(path)

        first_path = self.root / "first-distribution.xlsx"
        second_path = self.root / "second-distribution.xlsx"
        build(first_path, 1)
        build(second_path, 4)
        first = self.service.analyze(self.service.load_orders(first_path))
        second = self.service.analyze(self.service.load_orders(second_path))
        first_suggestion = first.suggest_slotting_thresholds({"A", "B", "C"}, 0.5)
        second_suggestion = second.suggest_slotting_thresholds({"A", "B", "C"}, 0.5)

        self.assertEqual(first_suggestion["method"], "empirical_relationship_pareto_knee")
        self.assertNotEqual(
            first_suggestion["minimum_shared_store_days"],
            second_suggestion["minimum_shared_store_days"],
        )
        self.assertGreater(second_suggestion["minimum_shared_store_days"], 1)
        self.assertGreater(first_suggestion["retained_weight_fraction"], 0)

    def test_export_writes_json_and_two_csv_files(self):
        dataset = self.service.load_orders(self.workbook_path)
        analysis = self.service.analyze(dataset)
        paths = self.service.export(
            analysis, self.root / "review.affinity.json",
            min_shared_store_days=1, top_per_sku=20,
        )
        self.assertTrue(all(path.exists() for path in paths))
        payload = json.loads(paths[0].read_text(encoding="utf-8"))
        self.assertEqual(payload["schema"], EXPORT_SCHEMA)
        self.assertEqual(payload["summary"]["line_order_events"], 6)
        self.assertEqual(payload["summary"]["store_day_groups"], 4)
        self.assertEqual(
            payload["filters"]["metric"],
            "cosine_similarity_of_binary_sku_presence_by_store_id_and_date",
        )
        self.assertTrue(payload["sku_store_frequencies"])
        self.assertTrue(payload["sku_relationships"])
        with paths[1].open(encoding="utf-8", newline="") as stream:
            store_rows = list(csv.DictReader(stream))
        with paths[2].open(encoding="utf-8", newline="") as stream:
            pair_rows = list(csv.DictReader(stream))
        self.assertEqual(len(store_rows), analysis.observed_pairs)
        self.assertTrue(any({row["sku_a"], row["sku_b"]} == {"A", "B"} for row in pair_rows))
        self.assertIn("shared_store_day_count", pair_rows[0])


if __name__ == "__main__":
    unittest.main()
