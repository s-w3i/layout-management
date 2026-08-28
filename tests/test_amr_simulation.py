from __future__ import annotations

import json
import math
from dataclasses import replace
from datetime import date, datetime, time
from pathlib import Path

import pytest
from openpyxl import Workbook

from amr_simulation.debugger import PlaybackController, _format_elapsed
from amr_simulation.coordination import (
    following_queue,
    reversed_passage,
    same_direction_following,
    wait_cycle,
)
from amr_simulation.conflict_solver import (
    has_extended_head_to_head,
    has_trivial_cycle,
    is_partial_solvable,
    partial_conflicts,
)
from amr_simulation.engine import simulate_day
from amr_simulation.run_simulation import _duration, main
from amr_simulation.inputs import (
    assign_workstations,
    load_workload,
    validate_inputs,
)
from amr_simulation.models import (
    MotionProfile,
    Rack,
    SimulationConfig,
    WorkloadTask,
    grid_name,
    grid_position,
)
from amr_simulation.results import summarize
from amr_simulation.routing import GridRouter, merge_straight_runs, motion_phases
from warehouse_layout.domain import GridProject, GridSpec, Marker
from warehouse_layout.rmf import RmfMapService


def config(spawns=("G0_0",), workstations=("WS",), overrides=None):
    return SimulationConfig(
        len(spawns), tuple(spawns), 0.0, tuple(workstations), overrides or {},
        4.0, 4.0, 10.0, MotionProfile(),
    )


def project():
    value = GridProject(GridSpec(width_m=4, length_m=2, spacing_m=1))
    value.markers = {
        (0, 0): Marker("rack", "RACK_0_0"),
        (2, 0): Marker("rack", "RACK_2_0"),
        (4, 0): Marker("rack", "RACK_4_0"),
        (2, 2): Marker("workstation", "WS"),
    }
    return value


def task(store, skus, release=0.0, picked_date=date(2024, 1, 1)):
    return WorkloadTask(
        f"{picked_date.isoformat()}/{store}", picked_date, release, store, dict(skus)
    )


def test_dram_pattern_matching_with_cycle_detection():
    assert has_extended_head_to_head((5, 6, 7), (7, 6, 5, 1))
    assert not has_extended_head_to_head((1, 2), (2, 3, 4))
    assert has_trivial_cycle(((1, 2), (2, 3), (3, 1)), 0)
    assert not has_trivial_cycle(((1, 2), (3, 4)), 0)
    assert not is_partial_solvable(((5, 6, 7), (7, 6, 5, 1)), 0)
    assert not is_partial_solvable(((1, 2), (2, 3), (3, 1)), 0)
    assert is_partial_solvable(((1, 2), (3, 4)), 0)
    head_to_head = partial_conflicts(((5, 6, 7), (7, 6, 5, 1)), 0)
    assert head_to_head[0].kind == "head_to_head"
    assert head_to_head[0].agent_indices == (0, 1)
    assert head_to_head[0].overlap_node == 6
    cycle = partial_conflicts(((1, 2), (2, 3), (3, 1)), 0)
    assert cycle[-1].kind == "cycle"
    assert cycle[-1].agent_indices == (0, 1, 2)
    cases = (
        ((1, 2, 3), (4, 5, 6)),
        ((6, 5, 9), (5, 6, 7)),
        ((1, 2), (2, 3, 4)),
        ((1, 2), (2, 3), (3, 1)),
    )
    for paths in cases:
        for index in range(len(paths)):
            assert bool(partial_conflicts(paths, index)) != is_partial_solvable(paths, index)


def test_dram_candidate_filter_preserves_conflicts():
    paths = ((1, 2, 3, 4), (4, 3, 2, 1), (9, 10), (10, 11, 9))
    candidates = tuple(
        index for index, path in enumerate(paths) if paths[0][0] in path[1:]
    )
    assert partial_conflicts(paths, 0, candidates) == partial_conflicts(paths, 0)


def test_dram_wait_configuration_is_snapshotted():
    selected = config()
    assert selected.dram_conflict_wait_seconds == 15.0
    assert selected.snapshot()["dram_conflict_wait_seconds"] == 15.0
    assert selected.loaded_priority_enabled
    assert selected.mutex_passage_enabled
    assert selected.corridor_coordination_enabled


def test_passage_following_queue_and_transitive_cycle_helpers():
    east = ((0, 0), (1, 0), (2, 0), (3, 0))
    west = tuple(reversed(east))
    leader = ((1, 0), (2, 0), (3, 0))
    assert reversed_passage(east, west) == east
    assert same_direction_following(east, leader, (1, 0))
    assert not same_direction_following(east, west, (1, 0))
    dependencies = {"B": {"A"}, "C": {"B"}, "A": set()}
    assert following_queue("A", dependencies) == ("A", "B", "C")
    assert wait_cycle({"A": {"B"}, "B": {"C"}, "C": {"A"}}) == ("A", "B", "C")


def test_workbook_grouping_ignores_time_preserves_line_counts_and_cache(tmp_path):
    source = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Time", "Store ID", "Item or SKU"])
    sheet.append([datetime(2024, 1, 1), time(8, 0), "S1", "A"])
    sheet.append([datetime(2024, 1, 1), time(8, 5), "S1", "A"])
    sheet.append([datetime(2024, 1, 1), time(8, 3), "S1", "B"])
    workbook.save(source)
    first = load_workload(source, tmp_path / "cache")
    second = load_workload(source, tmp_path / "cache")
    assert first.tasks[0].release_seconds == 0.0
    assert first.tasks[0].line_counts == {"A": 2, "B": 1}
    assert first.valid_rows == 3
    assert not first.cache_used and second.cache_used

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Store ID", "Item or SKU"])
    sheet.append([datetime(2024, 1, 2), "S2", "C"])
    workbook.save(source)
    changed = load_workload(source, tmp_path / "cache")
    assert not changed.cache_used and changed.tasks[0].store_id == "S2"


def test_workbook_rejects_bad_required_rows(tmp_path):
    source = tmp_path / "bad.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Store ID", "Item or SKU"])
    sheet.append([datetime(2024, 1, 1), "S1", None])
    workbook.save(source)
    try:
        load_workload(source, tmp_path / "cache", use_cache=False)
    except ValueError as exc:
        assert "row numbers: [2]" in str(exc)
    else:
        raise AssertionError("invalid row was accepted")


def test_balanced_fixed_workstation_assignment_and_override():
    tasks = [task("A", {"X": 5}), task("B", {"X": 4}), task("C", {"X": 3})]
    chosen = assign_workstations(tasks, config(workstations=("W1", "W2")))
    assert chosen == {"A": "W1", "B": "W2", "C": "W2"}
    overridden = assign_workstations(
        tasks, config(workstations=("W1", "W2"), overrides={"A": "W2"})
    )
    assert overridden["A"] == "W2" and overridden["B"] == "W1"


def test_directed_astar_rack_obstacle_and_motion_timing():
    value = GridProject(GridSpec(width_m=1, length_m=1, spacing_m=1))
    value.deleted_positions = {(0, 1), (1, 1)}
    value.markers = {
        (0, 0): Marker("rack", "R0"),
        (1, 0): Marker("rack", "R1"),
    }
    value.one_way_lanes = {((0, 0), (1, 0))}
    router = GridRouter(value)
    assert router.route((0, 0), (1, 0)) is not None
    assert router.route((1, 0), (0, 0)) is None
    assert router.reachable((0, 0), (1, 0))
    assert not router.reachable((1, 0), (0, 0))

    triangular = motion_phases(1.0, 1.5, 0.75)
    trapezoidal = motion_phases(6.0, 1.5, 0.75)
    assert [phase[0] for phase in triangular] == ["accel", "decel"]
    assert [phase[0] for phase in trapezoidal] == ["accel", "cruise", "decel"]
    runs = merge_straight_runs(project(), ((0, 0), (1, 0), (2, 0), (2, 1)))
    assert [(run.distance_m, round(math.degrees(run.heading))) for run in runs] == [(2.0, 0), (1.0, 90)]


def test_astar_prefers_fewer_turns_for_equal_distance():
    value = GridProject(GridSpec(width_m=3, length_m=2, spacing_m=1))
    router = GridRouter(value)
    edges = (
        ((0, 0), (1, 0)), ((1, 0), (1, 1)), ((1, 1), (2, 1)),
        ((2, 1), (2, 2)), ((2, 2), (3, 2)),
        ((0, 0), (0, 1)), ((0, 1), (0, 2)), ((0, 2), (1, 2)),
        ((1, 2), (2, 2)),
    )
    router.graph = {position: [] for position in router.positions}
    for start, end in edges:
        router.graph[start].append((end, 1.0))
    for neighbours in router.graph.values():
        neighbours.sort()
    route = router.route((0, 0), (3, 2))
    assert route.distance_m == 5.0
    assert route.positions == ((0, 0), (0, 1), (0, 2), (1, 2), (2, 2), (3, 2))
    assert len(merge_straight_runs(value, route.positions)) == 2


def test_greedy_rack_coverage_and_completion_at_jack_down():
    value, router = project(), GridRouter(project())
    racks = {
        "G2_0": Rack("G2_0", (2, 0), frozenset({"A", "B"})),
        "G4_0": Rack("G4_0", (4, 0), frozenset({"A"})),
    }
    result = simulate_day(
        value, router, racks, [task("S", {"A": 2, "B": 3}, 100.0)],
        {"S": "WS"}, config(), trace=True,
    )
    assert result.metrics["rack_presentations"] == 1
    assert result.metrics["completed_lines"] == 5
    assert result.jobs[0]["rack_id"] == "G2_0"
    jack_down = next(event for event in result.events if event["event"] == "jack_down_done")
    job = result.jobs[0]
    assert job["dispatch_time"] < job["rack_departure"] < job["rack_return"] < job["completion_time"]
    assert result.metrics["first_release_seconds"] == 0.0
    assert result.metrics["final_completion_seconds"] == jack_down["time_seconds"]
    assert result.metrics["makespan_seconds"] == jack_down["time_seconds"]

    planned = [event for event in result.events if event["event"] == "stage_path_planned"]
    assert [event["stage"] for event in planned] == ["to_pickup", "to_station", "return_rack"]
    event_times = {event["event"]: event["time_seconds"] for event in result.events}
    assert planned[0]["time_seconds"] == event_times["dispatch"]
    assert planned[1]["time_seconds"] == event_times["jack_up_done"]
    assert planned[2]["time_seconds"] == event_times["service_done"]


def test_parallel_jobs_replica_fallback_station_fifo_and_reservation():
    value = project()
    router = GridRouter(value)
    two_amrs = config(spawns=("G0_0", "G4_0"))
    racks = {
        "G2_0": Rack("G2_0", (2, 0), frozenset({"A"})),
        "G4_0": Rack("G4_0", (4, 0), frozenset({"A", "B"})),
    }
    result = simulate_day(
        value, router, racks,
        [task("S1", {"A": 1}), task("S2", {"A": 1, "B": 1})],
        {"S1": "WS", "S2": "WS"}, two_amrs, trace=True,
    )
    first_dispatches = [event for event in result.events if event["event"] == "dispatch"]
    assert len(first_dispatches) >= 2
    assert first_dispatches[0]["time_seconds"] == first_dispatches[1]["time_seconds"] == 0.0
    assert {first_dispatches[0]["rack_id"], first_dispatches[1]["rack_id"]} == {"G2_0", "G4_0"}
    starts = sorted(job["service_start"] for job in result.jobs)
    assert starts[1] >= starts[0] + two_amrs.service_seconds
    ordered = sorted(result.jobs, key=lambda job: job["service_start"])
    assert ordered[1]["station_admitted"] >= ordered[0]["service_start"] + two_amrs.service_seconds
    assert result.metrics["station_queue_time_seconds"] > 0
    assert len({job["station_queue_position"] for job in result.jobs}) == 2
    assert "G2_2" not in {job["station_queue_position"] for job in result.jobs}
    assert all(job["station_queue_enter"] <= job["station_admitted"] < job["station_arrival"] for job in result.jobs)
    assert result.metrics["reservation_conflicts"] > 0
    assert result.metrics["reservation_conflicts"] == (
        result.metrics["node_ownership_conflicts"]
        + result.metrics["dram_solver_conflicts"]
    )
    assert result.metrics["node_reservation_wait_seconds"] > 0

    owners = {
        (0, 0): "AMR_01",
        (4, 0): "AMR_02",
    }
    for event in result.events:
        if event["event"] == "reservation_granted":
            for node_name in event["nodes"]:
                node = grid_position(node_name)
                assert owners.get(node, event["amr_id"]) == event["amr_id"]
                owners[node] = event["amr_id"]
        elif event["event"] == "node_released":
            node = grid_position(event["node"])
            assert owners.pop(node) == event["amr_id"]

    same_rack = {"G2_0": Rack("G2_0", (2, 0), frozenset({"A"}))}
    sequential = simulate_day(
        value, GridRouter(value), same_rack,
        [task("S1", {"A": 1}), task("S2", {"A": 1})],
        {"S1": "WS", "S2": "WS"}, two_amrs, trace=True,
    )
    dispatches = [event["time_seconds"] for event in sequential.events if event["event"] == "dispatch"]
    completions = [event["time_seconds"] for event in sequential.events if event["event"] == "jack_down_done"]
    assert dispatches[1] == completions[0]


def test_workstation_queue_uses_arrival_time_not_task_sequence():
    value = project()
    racks = {
        "G2_0": Rack("G2_0", (2, 0), frozenset({"NEAR"})),
        "G4_0": Rack("G4_0", (4, 0), frozenset({"FAR"})),
    }
    result = simulate_day(
        value, GridRouter(value), racks,
        [task("S1", {"FAR": 1}), task("S2", {"NEAR": 1})],
        {"S1": "WS", "S2": "WS"},
        config(spawns=("G0_0", "G4_0")), trace=True,
    )
    by_task = {job["task_id"]: job for job in result.jobs}
    earlier_task = by_task["2024-01-01/S1"]
    later_task = by_task["2024-01-01/S2"]
    assert later_task["station_queue_enter"] < earlier_task["station_queue_enter"]
    assert later_task["service_start"] < earlier_task["service_start"]


def test_dynamic_tabu_route_and_continuous_node_crossings():
    value = project()
    router = GridRouter(value)
    direct = router.route((0, 1), (4, 1))
    detour = router.route((0, 1), (4, 1), frozenset({(2, 1)}))
    assert direct is not None and (2, 1) in direct.positions
    assert detour is not None and (2, 1) not in detour.positions
    edge_detour = router.route(
        (0, 1), (4, 1), frozenset(), frozenset({((1, 1), (2, 1))})
    )
    assert edge_detour is not None
    assert ((1, 1), (2, 1)) not in set(zip(edge_detour.positions, edge_detour.positions[1:]))

    route = router.route((0, 1), (2, 1))
    end, _heading, segments = router.motion(
        "J", "AMR", route, 0.0, 0.0, MotionProfile(), False
    )
    crossings = router.crossing_times(route, segments)
    assert [node for node, _time in crossings] == [(1, 1), (2, 1)]
    assert 0 < crossings[0][1] < crossings[1][1] == pytest.approx(end)


def test_daily_reset_batch_debug_parity_and_summary_math():
    value, racks = project(), {"G2_0": Rack("G2_0", (2, 0), frozenset({"A"}))}
    first = simulate_day(value, GridRouter(value), racks, [task("S", {"A": 2})], {"S": "WS"}, config())
    traced = simulate_day(value, GridRouter(value), racks, [task("S", {"A": 2})], {"S": "WS"}, config(), trace=True)
    assert first.metrics == traced.metrics
    assert traced.events and traced.motion_segments
    controller = PlaybackController.from_result(traced)
    controller.toggle()
    controller.advance(1.0)
    controller.next_event()
    controller.restart()
    assert controller.time == traced.metrics["first_release_seconds"]
    assert _format_elapsed(3661.25) == "01h 01m 01.25s"
    assert _duration(3661.25) == "01:01:01.25"

    second_task = task("S", {"A": 4}, picked_date=date(2024, 1, 2))
    second = simulate_day(value, GridRouter(value), racks, [second_task], {"S": "WS"}, config())
    summary = summarize("layout", [first, second])
    expected_weighted = (
        first.metrics["completed_lines"] + second.metrics["completed_lines"]
    ) / (first.metrics["makespan_hours"] + second.metrics["makespan_hours"])
    assert math.isclose(summary["weighted_throughput_lines_per_hour"], expected_weighted)


def test_coordination_switches_off_preserve_motion_results():
    value, racks = project(), {"G2_0": Rack("G2_0", (2, 0), frozenset({"A"}))}
    enabled = simulate_day(
        value, GridRouter(value), racks, [task("S", {"A": 1})], {"S": "WS"}, config()
    )
    disabled_config = replace(
        config(), loaded_priority_enabled=False, mutex_passage_enabled=False,
        corridor_coordination_enabled=False,
    )
    disabled = simulate_day(
        value, GridRouter(value), racks, [task("S", {"A": 1})],
        {"S": "WS"}, disabled_config,
    )
    keys = ("makespan_seconds", "travel_distance_m", "travel_time_seconds", "completed_lines")
    assert {key: enabled.metrics[key] for key in keys} == {
        key: disabled.metrics[key] for key in keys
    }
    assert enabled.jobs == disabled.jobs


def test_two_layout_hand_calculated_comparison():
    value = project()
    combined = {"G2_0": Rack("G2_0", (2, 0), frozenset({"A", "B"}))}
    split = {
        "G2_0": Rack("G2_0", (2, 0), frozenset({"A"})),
        "G4_0": Rack("G4_0", (4, 0), frozenset({"B"})),
    }
    workload = [task("S", {"A": 1, "B": 1})]
    one = simulate_day(value, GridRouter(value), combined, workload, {"S": "WS"}, config())
    two = simulate_day(value, GridRouter(value), split, workload, {"S": "WS"}, config())
    assert one.metrics["rack_presentations"] == 1
    assert two.metrics["rack_presentations"] == 2
    assert one.metrics["travel_distance_m"] == 6.0
    assert two.metrics["travel_distance_m"] == 16.0
    assert one.metrics["station_queue_time_seconds"] == two.metrics["station_queue_time_seconds"] == 0.0
    assert one.metrics["line_throughput_per_hour"] > two.metrics["line_throughput_per_hour"]


def test_cli_rejects_debug_without_one_date():
    try:
        main([
            "--mode", "debug", "--grid", "grid.json", "--orders", "orders.xlsx",
            "--config", "config.json", "--layout", "layout.json",
        ])
    except SystemExit as exc:
        assert exc.code == 2
    else:
        raise AssertionError("invalid debug arguments were accepted")


def test_amr_count_override_uses_configured_spawns():
    base = config(spawns=("G0_0", "G1_0", "G2_0"))
    selected = base.with_amr_count(2)
    assert selected.amr_count == 2
    assert selected.spawn_nodes == ("G0_0", "G1_0")

    with pytest.raises(ValueError, match="between 1 and 3"):
        base.with_amr_count(4)


def test_map1_regression_reports_42_inaccessible_racks():
    root = Path(__file__).resolve().parent.parent
    value = RmfMapService().load_project(root / "resources/map/map1.grid.json")
    router = GridRouter(value)
    racks = {
        grid_name(position): Rack(grid_name(position), position, frozenset())
        for position in router.rack_positions
    }
    selected_config = config(spawns=("G0_0",), workstations=tuple(router.workstations))
    report = validate_inputs(value, router, racks, [], selected_config, {})
    assert len([error for error in report.errors if error["code"] == "unreachable_rack"]) == 42
