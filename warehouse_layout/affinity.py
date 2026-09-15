"""Line-order SKU/store affinity analysis, caching, and export."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
from array import array
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Callable

import numpy as np
from openpyxl import load_workbook

from .config import DEFAULT_AFFINITY_CACHE


CACHE_SCHEMA = "sku_affinity_cache/v1"
EXPORT_SCHEMA = "sku_affinity_analysis/v2"
REQUIRED_COLUMNS = ("Date", "Store ID", "Item or SKU")

ProgressCallback = Callable[[int, int, str], None]
CancelCallback = Callable[[], bool]


class AffinityCancelledError(RuntimeError):
    """Raised when a caller cancels a workbook scan."""


def _excel_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            number = float(value)
        except (TypeError, ValueError):
            return None
        if not math.isfinite(number):
            return None
        return (datetime(1899, 12, 30) + timedelta(days=number)).date()
    try:
        return datetime.fromisoformat(str(value).strip()).date()
    except ValueError:
        return None


def _identifier(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass(slots=True)
class AffinityDataset:
    """Compact line-event dataset loaded from one workbook."""

    source_path: Path
    source_size: int
    source_mtime_ns: int
    worksheet: str
    dates: np.ndarray
    sku_indices: np.ndarray
    store_indices: np.ndarray
    skus: tuple[str, ...]
    stores: tuple[str, ...]
    valid_rows: int
    skipped_rows: int
    min_date: date
    max_date: date
    cache_used: bool = False


@dataclass(slots=True)
class AffinityAnalysis:
    """Date-filtered direct frequencies and store-day SKU affinities."""

    dataset: AffinityDataset
    start_date: date
    end_date: date
    frequency: np.ndarray
    sku_totals: np.ndarray
    store_totals: np.ndarray
    similarity: np.ndarray
    shared_store_days: np.ndarray
    sku_store_day_totals: np.ndarray
    store_day_count: int
    event_count: int
    observed_pairs: int

    def active_sku_indices(self) -> list[int]:
        return [int(index) for index in np.flatnonzero(self.sku_totals)]

    def top_skus(self, limit: int) -> list[int]:
        active = self.active_sku_indices()
        active.sort(
            key=lambda index: (
                -int(self.sku_totals[index]), self.dataset.skus[index].casefold()
            )
        )
        return active[: max(1, int(limit))]

    def resolve_sku(self, query: str) -> int | None:
        wanted = query.strip().casefold()
        if not wanted:
            top = self.top_skus(1)
            return top[0] if top else None
        active = self.active_sku_indices()
        exact = next(
            (index for index in active if self.dataset.skus[index].casefold() == wanted),
            None,
        )
        if exact is not None:
            return exact
        matches = [
            index for index in active
            if wanted in self.dataset.skus[index].casefold()
        ]
        if not matches:
            return None
        return min(
            matches,
            key=lambda index: (
                -int(self.sku_totals[index]), self.dataset.skus[index].casefold()
            ),
        )

    def common_store_labels(self, first: int, second: int, limit: int = 3) -> list[str]:
        first_counts = self.frequency[first]
        second_counts = self.frequency[second]
        common = np.flatnonzero((first_counts > 0) & (second_counts > 0))
        ranked = sorted(
            (int(index) for index in common),
            key=lambda index: (
                -min(int(first_counts[index]), int(second_counts[index])),
                -int(first_counts[index] + second_counts[index]),
                self.dataset.stores[index].casefold(),
            ),
        )
        return [
            f"{self.dataset.stores[index]} "
            f"({int(first_counts[index])}/{int(second_counts[index])})"
            for index in ranked[:limit]
        ]

    def related_skus(
        self, sku_index: int, min_shared_store_days: int = 3, limit: int = 12
    ) -> list[dict]:
        candidates = []
        for related in self.active_sku_indices():
            if related == sku_index:
                continue
            shared = int(self.shared_store_days[sku_index, related])
            score = float(self.similarity[sku_index, related])
            if shared < min_shared_store_days or score <= 0:
                continue
            candidates.append({
                "sku_index": related,
                "sku": self.dataset.skus[related],
                "affinity": score,
                "affinity_percent": score * 100.0,
                "shared_store_days": shared,
                "selected_orders": int(self.sku_totals[sku_index]),
                "related_orders": int(self.sku_totals[related]),
            })
        candidates.sort(
            key=lambda row: (
                -row["affinity"], -row["shared_store_days"],
                -row["related_orders"], row["sku"].casefold(),
            )
        )
        selected = candidates[: max(1, int(limit))]
        for row in selected:
            row["top_stores"] = self.common_store_labels(
                sku_index, int(row["sku_index"])
            )
        return selected

    def store_rows(self, sku_index: int) -> list[dict]:
        total = int(self.sku_totals[sku_index]) or 1
        rows = [
            {
                "store": self.dataset.stores[store_index],
                "orders": int(count),
                "share_percent": float(count) * 100.0 / total,
            }
            for store_index, count in enumerate(self.frequency[sku_index])
            if count > 0
        ]
        rows.sort(key=lambda row: (-row["orders"], row["store"].casefold()))
        return rows

    def suggest_slotting_thresholds(
        self, known_skus: set[str], affinity_weight: float
    ) -> dict:
        """Choose relationship thresholds from this dataset's Pareto knee.

        Candidate values are empirical quantiles of the observed relationships;
        no sample-specific support or score cutoff is embedded in the method.
        """
        if not 0.0 <= affinity_weight <= 1.0:
            raise ValueError("affinity weight must be between 0 and 1")
        indices = np.array(
            [
                index for index, sku in enumerate(self.dataset.skus)
                if sku in known_skus and self.sku_store_day_totals[index] > 0
            ],
            dtype=np.int32,
        )
        if len(indices) < 2:
            raise ValueError(
                "affinity workbook must contain at least two SKUs from the velocity CSV"
            )
        shared_matrix = self.shared_store_days[np.ix_(indices, indices)]
        score_matrix = self.similarity[np.ix_(indices, indices)]
        upper = np.triu_indices(len(indices), 1)
        shared = shared_matrix[upper].astype(np.int64)
        scores = score_matrix[upper].astype(np.float64)
        positive = (shared > 0) & (scores > 0)
        shared = shared[positive]
        scores = scores[positive]
        first = upper[0][positive]
        second = upper[1][positive]
        if not len(shared):
            raise ValueError(
                "affinity workbook has no shared store-day relationships among velocity SKUs"
            )
        mass = shared.astype(np.float64) * scores
        total_mass = float(mass.sum()) or 1.0
        related_sku_count = len(np.unique(np.concatenate((first, second))))

        # Sturges' rule determines search resolution from relationship count.
        candidate_count = max(2, int(math.ceil(math.log2(len(shared)) + 1)))
        quantiles = np.linspace(0.0, 1.0, candidate_count)
        support_candidates = np.unique(
            np.quantile(shared, quantiles, method="nearest").astype(np.int64)
        )
        score_candidates = np.unique(
            np.quantile(scores, quantiles, method="nearest").astype(np.float64)
        )
        support_bins = np.searchsorted(
            support_candidates, shared, side="right"
        ) - 1
        score_bins = np.searchsorted(score_candidates, scores, side="right") - 1
        grid_shape = (len(support_candidates), len(score_candidates))
        count_grid = np.zeros(grid_shape, dtype=np.int64)
        mass_grid = np.zeros(grid_shape, dtype=np.float64)
        np.add.at(count_grid, (support_bins, score_bins), 1)
        np.add.at(mass_grid, (support_bins, score_bins), mass)

        def reverse_cumulative(values):
            return values[::-1, ::-1].cumsum(axis=0).cumsum(axis=1)[::-1, ::-1]

        retained_counts = reverse_cumulative(count_grid)
        retained_mass = reverse_cumulative(mass_grid)

        # Count endpoint coverage for every threshold combination without
        # repeatedly scanning all relationships for every candidate pair.
        incident_skus = np.concatenate((first, second))
        incident_support = np.concatenate((support_bins, support_bins))
        incident_scores = np.concatenate((score_bins, score_bins))
        order = np.argsort(incident_skus, kind="stable")
        sorted_skus = incident_skus[order]
        boundaries = np.flatnonzero(
            np.r_[True, sorted_skus[1:] != sorted_skus[:-1], True]
        )
        coverage_grid = np.zeros(grid_shape, dtype=np.int64)
        for start, end in zip(boundaries[:-1], boundaries[1:]):
            positions = order[start:end]
            sku_grid = np.zeros(grid_shape, dtype=np.int8)
            sku_grid[
                incident_support[positions], incident_scores[positions]
            ] = 1
            coverage_grid += reverse_cumulative(sku_grid) > 0

        evaluations = []
        for support_position, support_threshold in enumerate(support_candidates):
            for score_position, score_threshold in enumerate(score_candidates):
                retained_count = int(
                    retained_counts[support_position, score_position]
                )
                if not retained_count:
                    continue
                edge_fraction = retained_count / len(shared)
                mass_fraction = float(
                    retained_mass[support_position, score_position]
                ) / total_mass
                coverage_fraction = float(
                    coverage_grid[support_position, score_position]
                ) / max(1, related_sku_count)
                # Distance to a sparse graph that still preserves useful
                # affinity mass and SKU coverage. Sparsity always retains a
                # minimum influence; otherwise 100% affinity selects nearly
                # every positive edge and destroys meaningful clusters.
                sparsity_weight = max(0.25, 1.0 - affinity_weight)
                preservation_weight = 1.0 - sparsity_weight
                loss = math.sqrt(
                    sparsity_weight * edge_fraction**2
                    + preservation_weight
                    * (
                        (1.0 - mass_fraction) ** 2
                        + (1.0 - coverage_fraction) ** 2
                    )
                    / 2.0
                )
                evaluations.append({
                    "minimum_shared_store_days": int(support_threshold),
                    "minimum_affinity_score": float(score_threshold),
                    "retained_relationship_count": retained_count,
                    "relationship_count": len(shared),
                    "retained_edge_fraction": edge_fraction,
                    "retained_weight_fraction": mass_fraction,
                    "sku_coverage_fraction": coverage_fraction,
                    "selection_loss": loss,
                })
        recommendation = min(
            evaluations,
            key=lambda row: (
                row["selection_loss"],
                -row["retained_weight_fraction"],
                row["retained_edge_fraction"],
                -row["sku_coverage_fraction"],
            ),
        )
        recommendation["method"] = "empirical_relationship_pareto_knee"
        recommendation["minimum_sparsity_weight"] = 0.25
        recommendation["candidate_support_count"] = len(support_candidates)
        recommendation["candidate_score_count"] = len(score_candidates)
        recommendation["active_affinity_sku_count"] = len(indices)
        return recommendation


class AffinityService:
    """Analyze line-level orders without changing slotting state."""

    def __init__(self, cache_dir: Path | None = None):
        self.cache_dir = Path(cache_dir or DEFAULT_AFFINITY_CACHE)

    @staticmethod
    def _fingerprint(path: Path) -> dict:
        resolved = path.expanduser().resolve()
        stat = resolved.stat()
        return {
            "source_path": str(resolved),
            "source_size": stat.st_size,
            "source_mtime_ns": stat.st_mtime_ns,
        }

    def _cache_path(self, path: Path) -> Path:
        digest = hashlib.sha256(
            str(path.expanduser().resolve()).encode("utf-8")
        ).hexdigest()[:20]
        return self.cache_dir / f"{digest}.npz"

    @staticmethod
    def _metadata_array(metadata: dict) -> np.ndarray:
        return np.frombuffer(
            json.dumps(metadata, separators=(",", ":")).encode("utf-8"),
            dtype=np.uint8,
        )

    @staticmethod
    def _decode_metadata(values: np.ndarray) -> dict:
        return json.loads(values.tobytes().decode("utf-8"))

    def _load_cache(self, path: Path, fingerprint: dict) -> AffinityDataset | None:
        cache_path = self._cache_path(path)
        if not cache_path.exists():
            return None
        try:
            with np.load(cache_path, allow_pickle=False) as cache:
                metadata = self._decode_metadata(cache["metadata"])
                if metadata.get("schema") != CACHE_SCHEMA:
                    return None
                if any(metadata.get(key) != value for key, value in fingerprint.items()):
                    return None
                dataset = AffinityDataset(
                    source_path=Path(metadata["source_path"]),
                    source_size=int(metadata["source_size"]),
                    source_mtime_ns=int(metadata["source_mtime_ns"]),
                    worksheet=str(metadata["worksheet"]),
                    dates=cache["dates"].astype(np.int32, copy=True),
                    sku_indices=cache["sku_indices"].astype(np.int32, copy=True),
                    store_indices=cache["store_indices"].astype(np.int32, copy=True),
                    skus=tuple(metadata["skus"]),
                    stores=tuple(metadata["stores"]),
                    valid_rows=int(metadata["valid_rows"]),
                    skipped_rows=int(metadata["skipped_rows"]),
                    min_date=date.fromisoformat(metadata["min_date"]),
                    max_date=date.fromisoformat(metadata["max_date"]),
                    cache_used=True,
                )
                if not (
                    len(dataset.dates)
                    == len(dataset.sku_indices)
                    == len(dataset.store_indices)
                    == dataset.valid_rows
                ):
                    raise ValueError("cache event arrays have inconsistent lengths")
                return dataset
        except Exception:
            try:
                cache_path.unlink()
            except OSError:
                pass
            return None

    def _write_cache(self, dataset: AffinityDataset) -> None:
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        cache_path = self._cache_path(dataset.source_path)
        temporary = cache_path.with_suffix(".tmp")
        metadata = {
            "schema": CACHE_SCHEMA,
            "source_path": str(dataset.source_path),
            "source_size": dataset.source_size,
            "source_mtime_ns": dataset.source_mtime_ns,
            "worksheet": dataset.worksheet,
            "skus": list(dataset.skus),
            "stores": list(dataset.stores),
            "valid_rows": dataset.valid_rows,
            "skipped_rows": dataset.skipped_rows,
            "min_date": dataset.min_date.isoformat(),
            "max_date": dataset.max_date.isoformat(),
        }
        try:
            with temporary.open("wb") as stream:
                np.savez_compressed(
                    stream,
                    metadata=self._metadata_array(metadata),
                    dates=dataset.dates,
                    sku_indices=dataset.sku_indices,
                    store_indices=dataset.store_indices,
                )
            os.replace(temporary, cache_path)
        finally:
            if temporary.exists():
                temporary.unlink()

    @staticmethod
    def _find_sheet(workbook) -> tuple[object, int, dict[str, int]]:
        required = set(REQUIRED_COLUMNS)
        for worksheet in workbook.worksheets:
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=1, max_row=min(50, worksheet.max_row), values_only=True),
                start=1,
            ):
                positions = {
                    str(value).strip(): index
                    for index, value in enumerate(values)
                    if value is not None
                }
                if required.issubset(positions):
                    return worksheet, row_number, positions
        raise ValueError(
            "No worksheet contains the required columns: "
            + ", ".join(REQUIRED_COLUMNS)
        )

    def load_orders(
        self,
        path: Path,
        progress: ProgressCallback | None = None,
        cancelled: CancelCallback | None = None,
        use_cache: bool = True,
    ) -> AffinityDataset:
        source = Path(path).expanduser().resolve()
        if not source.exists():
            raise ValueError(f"Order workbook not found: {source}")
        fingerprint = self._fingerprint(source)
        if use_cache:
            cached = self._load_cache(source, fingerprint)
            if cached is not None:
                if progress:
                    progress(cached.valid_rows, cached.valid_rows, "Loaded cached order events")
                return cached

        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            worksheet, header_row, positions = self._find_sheet(workbook)
            total = max(1, worksheet.max_row - header_row)
            date_index = positions["Date"]
            store_index = positions["Store ID"]
            sku_index = positions["Item or SKU"]
            dates = array("i")
            sku_indices = array("i")
            store_indices = array("i")
            sku_lookup: dict[str, int] = {}
            store_lookup: dict[str, int] = {}
            skipped = 0
            for processed, values in enumerate(
                worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=1
            ):
                if cancelled and cancelled():
                    raise AffinityCancelledError("Affinity analysis was cancelled")
                picked_date = _excel_date(
                    values[date_index] if date_index < len(values) else None
                )
                store = _identifier(
                    values[store_index] if store_index < len(values) else None
                )
                sku = _identifier(
                    values[sku_index] if sku_index < len(values) else None
                )
                if picked_date is None or not store or not sku:
                    skipped += 1
                else:
                    sku_id = sku_lookup.setdefault(sku, len(sku_lookup))
                    store_id = store_lookup.setdefault(store, len(store_lookup))
                    dates.append(picked_date.toordinal())
                    sku_indices.append(sku_id)
                    store_indices.append(store_id)
                if progress and (processed == 1 or processed % 5000 == 0):
                    progress(processed, total, f"Reading {worksheet.title}: {processed:,}/{total:,}")
            if not dates:
                raise ValueError("The workbook contains no valid SKU/store order rows")
            date_values = np.frombuffer(dates, dtype=np.int32).copy()
            dataset = AffinityDataset(
                source_path=source,
                source_size=int(fingerprint["source_size"]),
                source_mtime_ns=int(fingerprint["source_mtime_ns"]),
                worksheet=worksheet.title,
                dates=date_values,
                sku_indices=np.frombuffer(sku_indices, dtype=np.int32).copy(),
                store_indices=np.frombuffer(store_indices, dtype=np.int32).copy(),
                skus=tuple(sku_lookup),
                stores=tuple(store_lookup),
                valid_rows=len(date_values),
                skipped_rows=skipped,
                min_date=date.fromordinal(int(date_values.min())),
                max_date=date.fromordinal(int(date_values.max())),
                cache_used=False,
            )
            if progress:
                progress(total, total, "Building affinity cache")
            if use_cache:
                self._write_cache(dataset)
            return dataset
        finally:
            workbook.close()

    @staticmethod
    def analyze(
        dataset: AffinityDataset,
        start_date: date | None = None,
        end_date: date | None = None,
    ) -> AffinityAnalysis:
        start = start_date or dataset.min_date
        end = end_date or dataset.max_date
        if start > end:
            raise ValueError("Start date must be on or before end date")
        mask = (dataset.dates >= start.toordinal()) & (dataset.dates <= end.toordinal())
        event_count = int(np.count_nonzero(mask))
        if event_count == 0:
            raise ValueError("No valid order events fall within the selected date range")
        frequency = np.zeros((len(dataset.skus), len(dataset.stores)), dtype=np.int32)
        np.add.at(
            frequency,
            (dataset.sku_indices[mask], dataset.store_indices[mask]),
            1,
        )
        sku_totals = frequency.sum(axis=1, dtype=np.int64)
        store_totals = frequency.sum(axis=0, dtype=np.int64)
        # A fulfillment group is one Store ID on one Date. SKU presence is
        # binary inside the group, so duplicate lines do not inflate pair
        # affinity. Direct line frequency remains available in ``frequency``.
        event_dates = dataset.dates[mask]
        event_stores = dataset.store_indices[mask]
        event_skus = dataset.sku_indices[mask]
        order = np.lexsort((event_skus, event_stores, event_dates))
        grouped_dates = event_dates[order]
        grouped_stores = event_stores[order]
        grouped_skus = event_skus[order]
        unique_presence = np.ones(len(order), dtype=bool)
        unique_presence[1:] = (
            (grouped_dates[1:] != grouped_dates[:-1])
            | (grouped_stores[1:] != grouped_stores[:-1])
            | (grouped_skus[1:] != grouped_skus[:-1])
        )
        grouped_dates = grouped_dates[unique_presence]
        grouped_stores = grouped_stores[unique_presence]
        grouped_skus = grouped_skus[unique_presence]
        group_start = np.ones(len(grouped_skus), dtype=bool)
        group_start[1:] = (
            (grouped_dates[1:] != grouped_dates[:-1])
            | (grouped_stores[1:] != grouped_stores[:-1])
        )
        group_indices = np.cumsum(group_start, dtype=np.int32) - 1
        store_day_count = int(group_indices[-1]) + 1
        incidence = np.zeros(
            (len(dataset.skus), store_day_count), dtype=np.float32
        )
        incidence[grouped_skus, group_indices] = 1.0
        cooccurrence = np.rint(incidence @ incidence.T).astype(np.int32)
        store_day_totals = np.rint(incidence.sum(axis=1)).astype(np.int32)
        denominator = np.sqrt(
            np.outer(store_day_totals, store_day_totals).astype(np.float64)
        )
        similarity = np.divide(
            cooccurrence,
            denominator,
            out=np.zeros_like(denominator, dtype=np.float64),
            where=denominator > 0,
        )
        np.fill_diagonal(similarity, 0.0)
        np.fill_diagonal(cooccurrence, 0)
        return AffinityAnalysis(
            dataset=dataset,
            start_date=start,
            end_date=end,
            frequency=frequency,
            sku_totals=sku_totals,
            store_totals=store_totals,
            similarity=similarity,
            shared_store_days=cooccurrence,
            sku_store_day_totals=store_day_totals,
            store_day_count=store_day_count,
            event_count=event_count,
            observed_pairs=int(np.count_nonzero(frequency)),
        )

    @staticmethod
    def _export_paths(path: Path) -> tuple[Path, Path, Path]:
        destination = Path(path).expanduser()
        name = destination.name
        if name.endswith(".affinity.json"):
            base = destination.with_name(name[:-14])
        elif destination.suffix:
            base = destination.with_suffix("")
        else:
            base = destination
        return (
            base.with_name(base.name + ".affinity.json"),
            base.with_name(base.name + "_sku_store.csv"),
            base.with_name(base.name + "_sku_pairs.csv"),
        )

    def export(
        self,
        analysis: AffinityAnalysis,
        path: Path,
        min_shared_store_days: int = 3,
        top_per_sku: int = 20,
    ) -> tuple[Path, Path, Path]:
        json_path, store_path, pair_path = self._export_paths(path)
        json_path.parent.mkdir(parents=True, exist_ok=True)
        sparse_rows = []
        for sku_index, store_index in np.argwhere(analysis.frequency > 0):
            sparse_rows.append({
                "sku": analysis.dataset.skus[int(sku_index)],
                "store_id": analysis.dataset.stores[int(store_index)],
                "line_order_frequency": int(analysis.frequency[sku_index, store_index]),
            })
        pair_map: dict[tuple[int, int], dict] = {}
        for sku_index in analysis.active_sku_indices():
            for row in analysis.related_skus(
                sku_index,
                min_shared_store_days=min_shared_store_days,
                limit=top_per_sku,
            ):
                related = int(row["sku_index"])
                first, second = sorted((sku_index, related))
                key = (first, second)
                if key not in pair_map:
                    pair_map[key] = {
                        "sku_a": analysis.dataset.skus[first],
                        "sku_b": analysis.dataset.skus[second],
                        "affinity_percent": round(
                            float(analysis.similarity[first, second]) * 100.0, 6
                        ),
                        "shared_store_day_count": int(
                            analysis.shared_store_days[first, second]
                        ),
                        "sku_a_orders": int(analysis.sku_totals[first]),
                        "sku_b_orders": int(analysis.sku_totals[second]),
                        "top_contributing_stores": analysis.common_store_labels(first, second, 5),
                    }
        pair_rows = sorted(
            pair_map.values(),
            key=lambda row: (
                -row["affinity_percent"], -row["shared_store_day_count"],
                row["sku_a"].casefold(), row["sku_b"].casefold(),
            ),
        )
        payload = {
            "schema": EXPORT_SCHEMA,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "source": {
                "path": str(analysis.dataset.source_path),
                "size": analysis.dataset.source_size,
                "mtime_ns": analysis.dataset.source_mtime_ns,
                "worksheet": analysis.dataset.worksheet,
            },
            "filters": {
                "start_date": analysis.start_date.isoformat(),
                "end_date": analysis.end_date.isoformat(),
                "minimum_shared_store_days": int(min_shared_store_days),
                "top_relationships_per_sku": int(top_per_sku),
                "metric": "cosine_similarity_of_binary_sku_presence_by_store_id_and_date",
            },
            "summary": {
                "line_order_events": analysis.event_count,
                "active_skus": int(np.count_nonzero(analysis.sku_totals)),
                "active_stores": int(np.count_nonzero(analysis.store_totals)),
                "store_day_groups": analysis.store_day_count,
                "observed_sku_store_pairs": analysis.observed_pairs,
                "source_skipped_rows": analysis.dataset.skipped_rows,
            },
            "sku_totals": [
                {
                    "sku": analysis.dataset.skus[index],
                    "line_order_frequency": int(analysis.sku_totals[index]),
                }
                for index in analysis.active_sku_indices()
            ],
            "sku_store_frequencies": sparse_rows,
            "sku_relationships": pair_rows,
        }
        with json_path.open("w", encoding="utf-8") as stream:
            json.dump(payload, stream, indent=2)
            stream.write("\n")
        with store_path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(
                stream,
                fieldnames=("sku", "store_id", "line_order_frequency"),
                lineterminator="\n",
            )
            writer.writeheader()
            writer.writerows(sparse_rows)
        with pair_path.open("w", encoding="utf-8", newline="") as stream:
            fields = (
                "sku_a", "sku_b", "affinity_percent", "shared_store_day_count",
                "sku_a_orders", "sku_b_orders", "top_contributing_stores",
            )
            writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
            writer.writeheader()
            for row in pair_rows:
                writer.writerow({
                    **row,
                    "top_contributing_stores": "; ".join(row["top_contributing_stores"]),
                })
        return json_path, store_path, pair_path
