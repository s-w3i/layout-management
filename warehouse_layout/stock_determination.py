"""Reusable minimum-stock and buffer-stock calculations from order history."""

from __future__ import annotations

import csv
import math
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from itertools import permutations
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl import load_workbook


REQUIRED_COLUMNS = ("Date", "Item or SKU", "Quantity (in EA)")
CSV_FIELDS = (
    "sku",
    "observation_start",
    "observation_end",
    "observation_days",
    "total_demand_ea",
    "average_daily_demand_ea",
    "minimum_stock_days",
    "minimum_stock_ea",
    "buffer_stock_days",
    "minimum_buffer_stock_ea",
    "total_required_ea",
    "attribute_combination",
    "units_per_slot",
    "slots_per_unit",
    "occupied_level_span",
    "occupied_horizontal_slot_span",
    "required_slots",
    "required_racks",
    "rack_calculation_status",
)
COMBINATION_CSV_FIELDS = (
    "attribute_combination",
    "sku_count",
    "total_required_ea",
    "required_slots",
    "required_racks",
    "unresolved_skus",
)


def _date_value(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            number = float(value)
        except (TypeError, ValueError):
            return None
        if not math.isfinite(number):
            return None
        return (datetime(1899, 12, 30) + timedelta(days=number)).date()
    try:
        return datetime.fromisoformat(str(value).strip()).date()
    except ValueError:
        return None


def _sku_value(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_demand_workbook(path: Path) -> Iterator[tuple[str, date, float]]:
    """Yield valid ``(sku, date, nonnegative demand)`` rows from an XLSX file."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        positions = None
        for candidate in rows:
            names = {
                str(value).strip(): index
                for index, value in enumerate(candidate)
                if value is not None
            }
            if set(REQUIRED_COLUMNS).issubset(names):
                positions = names
                break
        if positions is None:
            raise ValueError(
                "Could not find a header row containing: "
                + ", ".join(REQUIRED_COLUMNS)
            )

        for values in rows:
            sku_index = positions["Item or SKU"]
            date_index = positions["Date"]
            quantity_index = positions["Quantity (in EA)"]
            sku = _sku_value(values[sku_index] if sku_index < len(values) else None)
            demand_date = _date_value(
                values[date_index] if date_index < len(values) else None
            )
            try:
                quantity = float(
                    values[quantity_index]
                    if quantity_index < len(values)
                    and values[quantity_index] not in (None, "")
                    else 0
                )
            except (TypeError, ValueError):
                quantity = 0.0
            if sku and demand_date is not None:
                yield sku, demand_date, (
                    quantity if math.isfinite(quantity) and quantity > 0 else 0.0
                )
    finally:
        workbook.close()


def calculate_stock_requirements(
    transactions: Iterable[tuple[str, date, float]],
    minimum_stock_days: int,
    buffer_stock_days: int,
) -> list[dict]:
    """Calculate whole-unit stock targets for every SKU in the input.

    Average daily demand uses the inclusive calendar span shared by the input.
    The combined minimum-plus-buffer target is rounded once so it is the least
    whole-unit quantity that covers the complete selected period.
    """
    if (
        isinstance(minimum_stock_days, bool)
        or not isinstance(minimum_stock_days, int)
    ):
        raise ValueError("minimum stock days must be a whole number")
    if (
        isinstance(buffer_stock_days, bool)
        or not isinstance(buffer_stock_days, int)
    ):
        raise ValueError("buffer stock days must be a whole number")
    if minimum_stock_days < 1:
        raise ValueError("minimum stock days must be at least 1")
    if buffer_stock_days < 0:
        raise ValueError("buffer stock days cannot be negative")

    totals: dict[str, float] = defaultdict(float)
    first_date = None
    last_date = None
    for sku, demand_date, quantity in transactions:
        sku = _sku_value(sku)
        if not sku or not isinstance(demand_date, date):
            continue
        try:
            quantity = float(quantity)
        except (TypeError, ValueError):
            quantity = 0.0
        totals[sku] += quantity if math.isfinite(quantity) and quantity > 0 else 0.0
        first_date = (
            demand_date if first_date is None else min(first_date, demand_date)
        )
        last_date = demand_date if last_date is None else max(last_date, demand_date)

    if first_date is None or last_date is None or not totals:
        raise ValueError("No valid SKU demand rows were found")
    observation_days = (last_date - first_date).days + 1
    results = []
    for sku in sorted(totals, key=lambda value: (value.casefold(), value)):
        total = totals[sku]
        daily = total / observation_days
        minimum = math.ceil(daily * minimum_stock_days)
        # Round the complete coverage target once. Rounding the minimum and
        # buffer independently can add an unnecessary unit and therefore an
        # unnecessary slot for low-volume SKUs.
        total_required = math.ceil(
            daily * (minimum_stock_days + buffer_stock_days)
        )
        buffer = total_required - minimum
        results.append({
            "sku": sku,
            "observation_start": first_date.isoformat(),
            "observation_end": last_date.isoformat(),
            "observation_days": observation_days,
            "total_demand_ea": round(total, 4),
            "average_daily_demand_ea": round(daily, 4),
            "minimum_stock_days": minimum_stock_days,
            "minimum_stock_ea": minimum,
            "buffer_stock_days": buffer_stock_days,
            "minimum_buffer_stock_ea": buffer,
            "total_required_ea": total_required,
        })
    return results


def determine_stock_requirements(
    workbook_path: Path,
    minimum_stock_days: int,
    buffer_stock_days: int,
) -> list[dict]:
    """Load an order workbook and return its per-SKU stock requirements."""
    return calculate_stock_requirements(
        read_demand_workbook(Path(workbook_path)),
        minimum_stock_days,
        buffer_stock_days,
    )


def _grid_units_per_slot(
    item_dimensions: tuple[float, float, float],
    slot_dimensions: tuple[float, float, float],
) -> int:
    """Return a physically constructible, rotation-aware box-grid capacity."""
    return max(
        (
            math.prod(
                math.floor(capacity / item + 1e-12)
                for item, capacity in zip(orientation, slot_dimensions)
            )
            for orientation in set(permutations(item_dimensions))
        ),
        default=0,
    )


def _quantity_rack_loads(
    total_required: int,
    required_slots: int,
    slots_per_unit: int,
    occupied_level_span: int,
    occupied_horizontal_slot_span: int,
    item_weight: float | None,
) -> list[tuple[int, int, float]]:
    """Materialize evenly divided slot loads for rack-capacity planning."""
    if total_required <= 0 or required_slots <= 0:
        return []
    footprint = max(1, slots_per_unit)
    load_count = max(1, math.ceil(required_slots / footprint))
    base_quantity, remainder = divmod(total_required, load_count)
    level_span = max(1, occupied_level_span)
    horizontal_span = max(1, occupied_horizontal_slot_span)
    return [
        (
            level_span,
            horizontal_span,
            (base_quantity + (index < remainder)) * (item_weight or 0.0),
        )
        for index in range(load_count)
    ]


def _rack_placements(
    occupied_mask: int,
    level_span: int,
    horizontal_span: int,
    levels_per_rack: int,
    slots_per_level: int,
) -> list[int]:
    """Return all distinct rack masks after placing one rectangular load."""
    results = []
    for level in range(levels_per_rack - level_span + 1):
        for slot in range(slots_per_level - horizontal_span + 1):
            placement = 0
            for row in range(level, level + level_span):
                for column in range(slot, slot + horizontal_span):
                    placement |= 1 << (row * slots_per_level + column)
            if occupied_mask & placement == 0:
                results.append(occupied_mask | placement)
    return sorted(set(results))


def _minimum_racks_for_loads(
    loads: Iterable[tuple[int, int, float]],
    levels_per_rack: int,
    slots_per_level: int,
    rack_max_weight: float | None,
) -> tuple[int, bool]:
    """Pack slot loads into racks, returning a feasible count and optimality flag.

    A lower-bound match proves the minimum immediately, which is the normal
    case when the default 12.5 kg slots are used with a 150 kg, 12-slot rack.
    Small nontrivial cases use exact branch-and-bound. Large custom cases retain
    the feasible best-fit result rather than understating rack demand.
    """
    normalized = [
        (int(level_span), int(horizontal_span), max(0.0, float(weight)))
        for level_span, horizontal_span, weight in loads
    ]
    if not normalized:
        return 0, True
    rack_slots = levels_per_rack * slots_per_level
    maximum_weight = math.inf if rack_max_weight is None else rack_max_weight
    for level_span, horizontal_span, weight in normalized:
        if (
            level_span < 1
            or horizontal_span < 1
            or level_span > levels_per_rack
            or horizontal_span > slots_per_level
            or weight > maximum_weight + 1e-9
        ):
            return 0, False

    if all(level == 1 and horizontal == 1 for level, horizontal, _ in normalized):
        if not math.isfinite(maximum_weight) or all(
            weight <= maximum_weight / rack_slots + 1e-9
            for _level, _horizontal, weight in normalized
        ):
            # Every possible full rack is weight-safe, so slot count alone is
            # both a constructive packing and a matching lower bound.
            return math.ceil(len(normalized) / rack_slots), True

    normalized.sort(
        key=lambda item: (
            item[0] * item[1] / rack_slots,
            0.0 if not math.isfinite(maximum_weight)
            else item[2] / maximum_weight,
            item[0], item[1], item[2],
        ),
        reverse=True,
    )

    # Best-fit decreasing supplies a feasible upper bound.
    greedy: list[tuple[int, float]] = []
    for level_span, horizontal_span, weight in normalized:
        candidates = []
        for index, (mask, used_weight) in enumerate(greedy):
            if used_weight + weight > maximum_weight + 1e-9:
                continue
            for updated_mask in _rack_placements(
                mask, level_span, horizontal_span,
                levels_per_rack, slots_per_level,
            ):
                candidates.append((
                    rack_slots - updated_mask.bit_count(),
                    maximum_weight - used_weight - weight,
                    index,
                    updated_mask,
                ))
        if candidates:
            _free_slots, _free_weight, index, updated_mask = min(candidates)
            greedy[index] = (updated_mask, greedy[index][1] + weight)
        else:
            placements = _rack_placements(
                0, level_span, horizontal_span,
                levels_per_rack, slots_per_level,
            )
            if not placements:
                return 0, False
            greedy.append((placements[0], weight))

    upper_bound = len(greedy)
    total_slots = sum(level * horizontal for level, horizontal, _ in normalized)
    lower_bound = math.ceil(total_slots / rack_slots)
    if math.isfinite(maximum_weight):
        lower_bound = max(
            lower_bound,
            math.ceil(sum(weight for _l, _h, weight in normalized) / maximum_weight),
        )

    identical_capacity_cache: dict[tuple[int, int], int] = {}

    def maximum_identical_per_rack(level_span: int, horizontal_span: int) -> int:
        cache_key = (level_span, horizontal_span)
        if cache_key in identical_capacity_cache:
            return identical_capacity_cache[cache_key]
        best = 0
        seen_masks: set[int] = set()

        def fill(mask: int, count: int) -> None:
            nonlocal best
            if mask in seen_masks:
                return
            seen_masks.add(mask)
            best = max(best, count)
            for updated_mask in _rack_placements(
                mask, level_span, horizontal_span,
                levels_per_rack, slots_per_level,
            ):
                if updated_mask > mask:
                    fill(updated_mask, count + 1)

        fill(0, 0)
        identical_capacity_cache[cache_key] = best
        return best

    # Repeated indivisible loads can make the aggregate area/weight lower
    # bounds too optimistic. Strengthen them using each identical load type's
    # maximum multiplicity in one rack.
    for (level_span, horizontal_span, weight), count in Counter(normalized).items():
        per_rack = maximum_identical_per_rack(level_span, horizontal_span)
        if math.isfinite(maximum_weight) and weight > 0:
            per_rack = min(per_rack, math.floor(maximum_weight / weight + 1e-12))
        if per_rack < 1:
            return 0, False
        lower_bound = max(lower_bound, math.ceil(count / per_rack))
    if lower_bound == upper_bound:
        return upper_bound, True

    # Exact search is intentionally limited to nontrivial small cases. The
    # result remains safe for larger custom datasets because the greedy count
    # is an actual packing, never merely a lower bound.
    if len(normalized) > 80:
        return upper_bound, False

    def can_pack(rack_count: int) -> bool:
        states = [(0, 0.0) for _ in range(rack_count)]

        def place(index: int) -> bool:
            if index == len(normalized):
                return True
            level_span, horizontal_span, weight = normalized[index]
            seen = set()
            for rack_index, (mask, used_weight) in enumerate(states):
                signature = (mask, round(used_weight, 9))
                if signature in seen:
                    continue
                seen.add(signature)
                if used_weight + weight > maximum_weight + 1e-9:
                    continue
                for updated_mask in _rack_placements(
                    mask, level_span, horizontal_span,
                    levels_per_rack, slots_per_level,
                ):
                    states[rack_index] = (updated_mask, used_weight + weight)
                    if place(index + 1):
                        return True
                    states[rack_index] = (mask, used_weight)
                if mask == 0:
                    break
            return False

        return place(0)

    for candidate in range(lower_bound, upper_bound):
        if can_pack(candidate):
            return candidate, True
    return upper_bound, True


def calculate_rack_requirements(
    stock_rows: Iterable[dict],
    sku_attributes: dict[str, dict],
    slot_dimensions: tuple[float, float, float],
    levels_per_rack: int,
    slots_per_level: int,
    combination_attributes: Iterable[str] = (),
    *,
    slot_max_weight: float | None = None,
    rack_max_weight: float | None = None,
) -> tuple[list[dict], list[dict]]:
    """Add slot/rack needs and summarize shared racks by attribute combination."""
    try:
        slot_dimensions = tuple(float(value) for value in slot_dimensions)
    except (TypeError, ValueError) as exc:
        raise ValueError("slot length, width, and height must be numeric") from exc
    if len(slot_dimensions) != 3 or any(
        not math.isfinite(value) or value <= 0 for value in slot_dimensions
    ):
        raise ValueError("slot length, width, and height must be greater than zero")
    if (
        isinstance(levels_per_rack, bool)
        or not isinstance(levels_per_rack, int)
        or levels_per_rack < 1
    ):
        raise ValueError("rack levels must be a whole number of at least 1")
    if (
        isinstance(slots_per_level, bool)
        or not isinstance(slots_per_level, int)
        or slots_per_level < 1
    ):
        raise ValueError("slots per level must be a whole number of at least 1")
    for label, value in (
        ("slot maximum weight", slot_max_weight),
        ("rack maximum weight", rack_max_weight),
    ):
        if value is None:
            continue
        try:
            normalized = float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{label} must be numeric") from exc
        if not math.isfinite(normalized) or normalized <= 0:
            raise ValueError(f"{label} must be greater than zero")
        if label.startswith("slot"):
            slot_max_weight = normalized
        else:
            rack_max_weight = normalized

    combination_attributes = tuple(dict.fromkeys(combination_attributes))
    enriched = []
    grouped: dict[tuple, dict] = {}
    for source_row in stock_rows:
        row = dict(source_row)
        sku = str(row.get("sku", "")).strip()
        attributes = dict(sku_attributes.get(sku) or {})
        signature = tuple(attributes.get(key) for key in combination_attributes)
        combination_text = " · ".join(
            f"{key}="
            f"{'T' if value is True else 'F' if value is False else '?'}"
            for key, value in zip(combination_attributes, signature)
        ) or "all SKUs"
        row["attribute_combination"] = combination_text
        status = "CALCULATED"
        item_dimensions = []
        for key in ("max_item_length", "max_item_width", "max_item_height"):
            try:
                value = float(attributes.get(key))
            except (TypeError, ValueError):
                value = 0.0
            if not math.isfinite(value) or value <= 0:
                status = "MISSING_DIMENSIONS"
            item_dimensions.append(value)

        try:
            item_weight = float(attributes.get("max_item_weight"))
        except (TypeError, ValueError):
            item_weight = 0.0
        weight_known = math.isfinite(item_weight) and item_weight > 0
        if not weight_known:
            item_weight = 0.0

        dimensional_units_per_slot = 0
        weight_units_per_slot = 0
        units_per_slot = 0
        slots_per_unit = 0
        occupied_level_span = 0
        occupied_horizontal_slot_span = 0
        if status == "MISSING_DIMENSIONS":
            # Keep the rack estimate complete without inventing dimensions.
            # One unknown unit per slot is conservative and remains visibly
            # marked as unverified in the exported/UI status.
            dimensional_units_per_slot = 1
            status = "CALCULATED_UNVERIFIED_DIMENSIONS"
        elif status == "CALCULATED":
            dimensional_units_per_slot = _grid_units_per_slot(
                tuple(item_dimensions), slot_dimensions
            )
            if dimensional_units_per_slot < 1:
                footprints = []
                for orientation in set(permutations(item_dimensions)):
                    if orientation[0] > slot_dimensions[0] + 1e-12:
                        continue
                    horizontal = max(
                        1, math.ceil(orientation[1] / slot_dimensions[1] - 1e-12)
                    )
                    vertical = max(
                        1, math.ceil(orientation[2] / slot_dimensions[2] - 1e-12)
                    )
                    if horizontal <= slots_per_level and vertical <= levels_per_rack:
                        footprints.append(
                            (vertical * horizontal, vertical, horizontal)
                        )
                if footprints:
                    (
                        slots_per_unit,
                        occupied_level_span,
                        occupied_horizontal_slot_span,
                    ) = min(footprints)
                    status = "CALCULATED_MULTI_SLOT"
                else:
                    status = "ITEM_DOES_NOT_FIT"
        total_required = max(0, int(math.ceil(float(row["total_required_ea"]))))
        if slot_max_weight is not None:
            weight_units_per_slot = (
                max(1, math.floor(slot_max_weight / item_weight))
                if weight_known else 1
            )
        if dimensional_units_per_slot:
            units_per_slot = dimensional_units_per_slot
            if weight_units_per_slot:
                units_per_slot = min(units_per_slot, weight_units_per_slot)
            required_slots = math.ceil(total_required / units_per_slot)
            slots_per_unit = 1
            occupied_level_span = 1
            occupied_horizontal_slot_span = 1
        elif slots_per_unit:
            required_slots = total_required * slots_per_unit
        else:
            required_slots = None
        rack_loads = (
            _quantity_rack_loads(
                total_required,
                required_slots,
                slots_per_unit,
                occupied_level_span,
                occupied_horizontal_slot_span,
                item_weight if weight_known else None,
            )
            if required_slots is not None else []
        )
        if required_slots is None:
            required_racks = None
            rack_minimum_proven = False
        else:
            required_racks, rack_minimum_proven = _minimum_racks_for_loads(
                rack_loads,
                levels_per_rack,
                slots_per_level,
                rack_max_weight,
            )
        if not weight_known and (
            slot_max_weight is not None or rack_max_weight is not None
        ):
            status = (
                "CALCULATED_UNVERIFIED_PHYSICAL"
                if status.startswith("CALCULATED") else status
            )
        elif required_slots is not None and not rack_minimum_proven:
            status = (
                "CALCULATED_FEASIBLE_RACK_UPPER_BOUND"
                if required_racks else "ITEM_DOES_NOT_FIT"
            )
        row.update({
            "units_per_slot": units_per_slot or "",
            "slots_per_unit": slots_per_unit or "",
            "occupied_level_span": occupied_level_span or "",
            "occupied_horizontal_slot_span": (
                occupied_horizontal_slot_span or ""
            ),
            "required_slots": required_slots if required_slots is not None else "",
            "required_racks": required_racks if required_racks is not None else "",
            "rack_calculation_status": status,
        })
        enriched.append(row)

        group = grouped.setdefault(signature, {
            "attributes": dict(zip(combination_attributes, signature)),
            "attribute_combination": combination_text,
            "sku_count": 0,
            "total_required_ea": 0,
            "_rack_loads": [],
            "required_slots": 0,
            "unresolved_skus": 0,
            "_unverified_weight_skus": 0,
        })
        group["sku_count"] += 1
        group["total_required_ea"] += total_required
        if not weight_known:
            group["_unverified_weight_skus"] += 1
        group["_rack_loads"].extend(rack_loads)
        if required_slots is None:
            group["unresolved_skus"] += 1
        else:
            group["required_slots"] += required_slots

    combinations = []
    for group in grouped.values():
        group["required_racks"], minimum_proven = _minimum_racks_for_loads(
            group["_rack_loads"],
            levels_per_rack,
            slots_per_level,
            rack_max_weight,
        )
        if not minimum_proven:
            group["unresolved_skus"] += 1
        if rack_max_weight is not None:
            group["unresolved_skus"] += group["_unverified_weight_skus"]
        group.pop("_rack_loads")
        group.pop("_unverified_weight_skus")
        combinations.append(group)
    combinations.sort(key=lambda row: row["attribute_combination"])
    return enriched, combinations


def write_stock_requirements_csv(rows: Iterable[dict], path: Path) -> Path:
    """Write calculated stock targets using the stable public CSV schema."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=CSV_FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(
            {field: row.get(field, "") for field in CSV_FIELDS} for row in rows
        )
    return path


def write_attribute_combination_csv(rows: Iterable[dict], path: Path) -> Path:
    """Write rack totals grouped by the selected SKU attribute combination."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(
            stream, fieldnames=COMBINATION_CSV_FIELDS, lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(
            {field: row.get(field, "") for field in COMBINATION_CSV_FIELDS}
            for row in rows
        )
    return path
